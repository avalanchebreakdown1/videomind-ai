import json
import os
from pathlib import Path
import tempfile


def write_text(path: Path, text: str):
    """Atomic replacement within a run; no half-written final artifact."""
    name = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                         prefix=".tmp-", delete=False) as handle:
            name = handle.name
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(name, path)
    finally:
        if name and Path(name).exists():
            Path(name).unlink()


def write_json(path: Path, value):
    write_text(path, json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def export_rewrite_docx(path: Path, document: dict, attribution: dict):
    """Clean educational notes: audit details stay in their own JSON artifacts."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    page = doc.sections[0]
    page.top_margin = page.bottom_margin = Inches(.8)
    page.left_margin = page.right_margin = Inches(.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("17365D")
    doc.core_properties.title = document["title"]
    doc.core_properties.author = "VideoMind AI"
    page.header.paragraphs[0].text = "VideoMind AI | Technical study notes"
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    page.footer.paragraphs[0]._p.append(field)
    doc.add_heading(document["title"], 0)
    doc.add_paragraph("Technical study notes", "Subtitle")
    for section in document["sections"]:
        doc.add_heading(section["heading"], 1)
        for block in section["blocks"]:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block["text"])
            if block["kind"] == "code":
                run.font.name, run.font.size = "Consolas", Pt(9)
                paragraph.paragraph_format.line_spacing = 1
    doc.add_heading("Source attribution", 1)
    doc.add_paragraph(attribution["title"])
    if attribution.get("creator"):
        doc.add_paragraph("Creator: " + attribution["creator"])
    if attribution.get("url"):
        doc.add_paragraph(attribution["url"])
    doc.add_paragraph("Adapted from the source transcript with contextual editorial corrections. "
                      "This document has not undergone independent factual verification.")
    name = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".tmp-rewrite-docx-", delete=False) as handle:
            name = handle.name
        doc.save(name)
        os.replace(name, path)
    finally:
        if name and Path(name).exists():
            Path(name).unlink()


def export_docx(path: Path, document: dict, knowledge: dict, evidence: list, quality: dict, attribution: dict):
    """Word export is deterministic rendering; model text cannot control document XML."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("17365D")
    doc.core_properties.title = document["title"]
    doc.core_properties.author = "VideoMind AI"
    doc.core_properties.subject = "Technical education — source-grounded draft"
    section.header.paragraphs[0].text = "VideoMind AI | Technical education"
    footer = section.footer.paragraphs[0]
    footer.text = "Source-grounded draft • Page "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_heading(document["title"], 0)
    doc.add_paragraph("Technical educational explanation", "Subtitle")
    status = doc.add_paragraph()
    status.add_run("REQUIRES REVIEW" if quality["state"] == "needs_review" else "SOURCE-CONSISTENCY CHECKS PASSED").bold = True
    doc.add_paragraph("AI-generated source-grounded draft. Scientific claims have not been independently fact-verified. "
                      "Review the source, qualifications and flagged items before using this document for teaching or publication.")
    doc.add_paragraph("Source [S1]: " + attribution["title"])
    if attribution.get("creator"):
        doc.add_paragraph("Creator: " + attribution["creator"])
    if attribution.get("url"):
        doc.add_paragraph("Source URL: " + attribution["url"])
    gate = document.get("content_gate")
    if gate and gate["counts"]["withheld"]:
        doc.add_paragraph(f"[REVIEW REQUIRED] {gate['counts']['withheld']} candidate paragraphs were withheld by the content gate. "
                          "This is an incomplete reading copy, not a complete account of the lecture. Full candidates and reasons are retained separately.")
    if gate and gate["state"] == "no_content_passed":
        doc.add_paragraph("No candidate paragraphs passed the content gate. A corrected educational explanation requires further source review.")
    concepts = {item["id"]: item for item in knowledge["concepts"]}
    for item in document["sections"]:
        doc.add_heading(item["heading"], 1)
        for paragraph in item["paragraphs"]:
            flags = [issue for issue in quality["issues"] if issue["item_id"] == paragraph["id"]]
            if flags:
                warning = doc.add_paragraph()
                run = warning.add_run("[REVIEW REQUIRED: " + "; ".join(sorted({issue["code"] for issue in flags})) + "]")
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("9C2929")
            refs = sorted({citation["evidence_id"] for key in paragraph["concept_ids"] for citation in concepts[key]["citations"]})
            doc.add_paragraph(paragraph["text"] + " [S1: " + ", ".join(refs) + "]")
    doc.add_heading("Source attribution and evidence", 1)
    doc.add_paragraph("[S1] " + attribution["title"] + (" — " + attribution["url"] if attribution.get("url") else " — local transcript (no public URL supplied)."))
    doc.add_paragraph("The explanation was synthesized from extracted concepts. Evidence locations identify source support, "
                      "not external scientific validation. Technical notation is retained as Unicode text.")
    used_concepts = {key for item in document["sections"] for paragraph in item["paragraphs"] for key in paragraph["concept_ids"]}
    cited = {citation["evidence_id"] for key in used_concepts for citation in concepts[key]["citations"]}
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, title in zip(table.rows[0].cells, ("Evidence", "Source time (seconds)", "Audio review status")):
        cell.text = title
    for item in evidence:
        if item["id"] in cited:
            cells = table.add_row().cells
            cells[0].text = item["id"]
            cells[1].text = f"{item['start']:.2f}–{item['end']:.2f}"
            cells[2].text = ", ".join(item["verification_statuses"])
    doc.add_heading("Content integrity review", 1)
    if gate:
        from collections import Counter
        doc.add_paragraph("Full findings and withheld wording remain in quality.json and content_gate.json; "
                          "they are not reproduced as educational content.")
        for code, count in sorted(Counter(issue["code"] for issue in quality["issues"]).items()):
            doc.add_paragraph(f"{code}: {count}", "List Bullet")
    else:
        for issue in quality["issues"]:
            doc.add_paragraph(f"{issue['item_id']} — {issue['code']}: {issue['detail']}", "List Bullet")
    if not quality["issues"]:
        doc.add_paragraph("No issues detected by the implemented source-consistency checks.")
    for limitation in quality["limitations"]:
        doc.add_paragraph(limitation, "List Bullet")
    name = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".tmp-docx-", delete=False) as handle:
            name = handle.name
        doc.save(name)
        os.replace(name, path)
    finally:
        if name and Path(name).exists():
            Path(name).unlink()


def export_reference_docx(path: Path, document: dict, references: dict, attribution: dict):
    """Render a separately edited reference-backed guide. Does not verify facts itself."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for section in document["sections"]:
        for item in [*section["paragraphs"], *section.get("examples", [])]:
            if not item.get("references") or not set(item["references"]) <= set(references):
                raise ValueError("Every edited paragraph/example requires known reference IDs")
    doc = Document()
    page = doc.sections[0]
    page.top_margin = page.bottom_margin = Inches(.8)
    page.left_margin = page.right_margin = Inches(.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for style in ("Heading 1", "Heading 2"):
        doc.styles[style].font.color.rgb = RGBColor.from_string("17365D")
    doc.core_properties.title = document["title"]
    doc.core_properties.author = "VideoMind AI — editorial revision"
    page.header.paragraphs[0].text = "VideoMind AI | Reference-checked study guide"
    footer = page.footer.paragraphs[0]
    footer.text = "Editorial revision • Page "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_heading(document["title"], 0)
    doc.add_paragraph(document["subtitle"], "Subtitle")
    doc.add_paragraph(document["scope"])
    for section in document["sections"]:
        doc.add_heading(section["heading"], 1)
        for paragraph in section["paragraphs"]:
            doc.add_paragraph(paragraph["text"] + " [" + ", ".join(paragraph["references"]) + "]")
        for example in section.get("examples", []):
            doc.add_paragraph(example["title"], "Heading 2")
            code = doc.add_paragraph()
            run = code.add_run(example["code"])
            run.font.name, run.font.size = "Consolas", Pt(9)
            doc.add_paragraph("Expected output: " + example["output"].strip())
            doc.add_paragraph("Language references: " + ", ".join(example["references"]))
    doc.add_heading("References and attribution", 1)
    doc.add_paragraph("Lecture topic source: " + attribution["title"] + " — " + attribution["url"])
    if attribution.get("creator"):
        doc.add_paragraph("Lecture creator: " + attribution["creator"])
    doc.add_paragraph("Technical statements in this revision are supported by the numbered official references below. "
                      "Examples are new teaching examples. This guide does not certify the accuracy or completeness of the original audio transcript.")
    for key, reference in references.items():
        doc.add_paragraph(f"[{key}] {reference['title']}\n{reference['url']}")
    doc.add_paragraph("References consulted: " + document["reference_check_date"] + ". " + document["review_limitations"])
    name = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".tmp-reference-docx-", delete=False) as handle:
            name = handle.name
        doc.save(name)
        os.replace(name, path)
    finally:
        if name and Path(name).exists():
            Path(name).unlink()
