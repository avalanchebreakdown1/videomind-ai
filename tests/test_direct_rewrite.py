"""Direct rewriting: real Word exports, controlled repairs, failure and API contracts."""
from copy import deepcopy
from dataclasses import replace
import io
import json

from docx import Document
from fastapi.testclient import TestClient
import pytest

from test_writing import source_dir, MeasurementASR
from videomind.api import create_app
from videomind.config import Settings
from videomind.ingestion_config import IngestionSettings
from videomind.services.direct_rewrite_service import transcript_chunks, unsupported_code_tokens, normalize_blocks, RewriteBlock, clean_reading_document
from videomind.services.ollama_service import OllamaError
from videomind.services.verification_service import sha256
from videomind.services.writing_pipeline import generate_document
from videomind.writing_config import WritingSettings


class DirectModel:
    def __init__(self, mode="good"):
        self.mode, self.requests = mode, []

    def metadata(self):
        return {"model": "controlled", "digest": "unchanged"}

    def chat(self, request):
        self.requests.append(deepcopy(request))
        assert request["format"]["title"] == "DirectRewrite"
        assert "recognisable ASR" in request["messages"][0]["content"]
        data = json.loads(request["messages"][1]["content"].split("\nUNTRUSTED_DATA_JSON:\n")[1])
        assert "CURRENT_TRANSCRIPT" in data
        if self.mode == "unavailable":
            raise OllamaError("Controlled offline failure")
        result = {"sections": [{"heading": "Optical measurements",
            "explanation": "The detector receives .5 mW over an area of 2 cm² under the specified conditions.",
            "code_examples": ['if True:\n    print("Hello friends")']}], "corrections": []}
        if self.mode == "repair":
            result["sections"][0]["explanation"] = "Python is a high-level programming language."
            result["corrections"] = [{"original": "Paitan", "replacement": "Python",
                                       "reason": "The surrounding lesson explicitly discusses Python programming."}]
        if self.mode == "empty":
            result["sections"] = []
        if self.mode == "invalid" or self.mode == "retry" and len(self.requests) == 1:
            result = {"wrong_schema": True}
        if self.mode == "truncated":
            return {"done": True, "done_reason": "length", "message": {"content": json.dumps(result)}}
        return {"done": True, "done_reason": "stop", "message": {"content": json.dumps(result)}}


def test_direct_context_rewrite_and_clean_word_preserve_source(source_dir):
    raw = source_dir / "verified_transcript.json"
    data = json.loads(raw.read_text())
    added = ' Paitan is a programming language. Python uses dynamic typing. Example: if True: print("Hello friends").'
    data["text"] += added
    data["words"].append({"text": added,
                          "start": 20, "end": 25, "probability": .5,
                          "verification_status": "not_checked", "audio_reviewed": False})
    # No saved preparation copy for this deliberately modified fixture.
    (source_dir / "document_source.json").unlink()
    raw.write_text(json.dumps(data))
    before = {p.name: sha256(p) for p in source_dir.iterdir()}
    model = DirectModel("repair")
    result = generate_document(source_dir, source_dir.parent / "direct", WritingSettings(), model,
                               title="Programming", url="https://example.org/lecture")
    assert len(model.requests) == 1
    word = "\n".join(p.text for p in Document(result).paragraphs)
    assert "Python is a high-level programming language." in word
    assert "Hello friends" in word  # Keep a greeting that is part of a code example.
    assert "[uncertain" not in word and "REVIEW REQUIRED" not in word
    assert "https://example.org/lecture" in word
    audit = json.loads(result.with_name("rewrite_audit.json").read_text())
    assert audit["corrections"][0]["original_found_in_source"]
    assert audit["corrections"][0]["replacement"] == "Python"
    assert not result.with_name("knowledge.json").exists()
    assert not result.with_name("content_gate.json").exists()
    assert {p.name: sha256(p) for p in source_dir.iterdir()} == before
    assert json.loads(result.with_name("integrity.json").read_text())["source_artifacts_unchanged"]


def test_direct_preserves_notation_and_code(source_dir):
    raw = source_dir / "verified_transcript.json"
    data = json.loads(raw.read_text())
    added = ' Example: if True: print("Hello friends").'
    data["text"] += added
    data["words"].append({"text": added, "start": 20, "end": 25, "probability": .99,
                          "verification_status": "not_checked", "audio_reviewed": False})
    (source_dir / "document_source.json").unlink()
    raw.write_text(json.dumps(data))
    result = generate_document(source_dir, source_dir.parent / "direct", WritingSettings(), DirectModel())
    word = Document(result)
    text = "\n".join(p.text for p in word.paragraphs)
    assert ".5 mW" in text and "2 cm²" in text
    code = next(p for p in word.paragraphs if p.text.startswith("if True:"))
    assert code.text == 'if True:\n    print("Hello friends")'
    assert code.runs[0].font.name == "Consolas"


def test_direct_excludes_unsourced_generated_code_and_audits_it(source_dir):
    result = generate_document(source_dir, source_dir.parent / "direct", WritingSettings(), DirectModel())
    text = "\n".join(p.text for p in Document(result).paragraphs)
    assert 'print("Hello friends")' not in text
    excluded = json.loads(result.with_name("rewrite_audit.json").read_text())["excluded_blocks"]
    assert excluded[0]["new_tokens"] == ["print"]
    assert 'print("Hello friends")' in excluded[0]["text"]


def test_import_guard_requires_an_actual_source_import():
    assert unsupported_code_tokens("import python", "Python is a programming language.") == ["unsourced import: python"]
    assert unsupported_code_tokens("import math", "Use import math to load the math module.") == []


def test_code_guard_accepts_spoken_numbers_local_names_and_comments():
    code = "a = 10\nb = 5\nresult = a + b  # Addition\nprint(result)"
    assert unsupported_code_tokens(code, "Create variables, one will be 10 and one five. Print the sum.") == []
    assert unsupported_code_tokens("item_name = 'orange'\nprint(item_name)", "Assign a variable item name with an underscore to orange and print it.") == []


def test_code_guard_accepts_reconstructed_function_but_not_new_calls_or_literals():
    assert unsupported_code_tokens("def greet_user(name):\n    print(name)\ngreet_user('Mario')", "Define a function to print a name and call it with Mario.") == []
    assert "dangerous" in unsupported_code_tokens("result = dangerous()", "Assign a variable and print it.")
    assert "alice" in unsupported_code_tokens("names = ['Alice']", "Assign a variable containing a list of names.")


def test_code_guard_recognizes_spoken_identifiers_and_conversion_builtin():
    assert unsupported_code_tokens("print(item_name)", "Print the item name.") == []
    assert unsupported_code_tokens("print(str(number))", "Print number. Convert the integer to a string.") == []


def test_code_guard_handles_fstring_syntax_and_contextual_conversion_exception():
    code = "def greet_user(name):\n    print(f'Hey there, {name}!')\ngreet_user('Mario')"
    assert unsupported_code_tokens(code, "Define a function to print Hey there with a name. Call it with Mario.") == []
    code = "try:\n    print(int(input()))\nexcept ValueError:\n    pass"
    assert unsupported_code_tokens(code, "Try to convert user input to an integer, print it, and handle an exception.") == []


def test_fenced_code_inside_a_paragraph_is_checked_as_code():
    parts = normalize_blocks(RewriteBlock(kind="paragraph", text="Explanation.\n```python\nimport python\n```\nMore prose."))
    assert [p.kind for p in parts] == ["paragraph", "code", "paragraph"]
    assert unsupported_code_tokens(parts[1].text, "Python is a language.") == ["unsourced import: python"]
    assert parts[0].text == "Explanation." and parts[2].text == "More prose."


def test_reading_cleanup_removes_actual_classroom_failure_preserving_technical_example():
    data = {"title": "Lesson", "sections": [{"heading": "Python", "blocks": [
        {"kind": "paragraph", "text": "Hello friends. Python is easy to learn. The function returns .5 mW. It is recommended to memorize the name and image of the creator.", "source_chunk_ids": ["T0"]},
        {"kind": "code", "text": 'print("Hello friends")', "source_chunk_ids": ["T0"]}]}]}
    original = deepcopy(data)
    reading, exclusions = clean_reading_document(data)
    assert reading["sections"][0]["blocks"][0]["text"] == "The function returns .5 mW."
    assert reading["sections"][0]["blocks"][1]["text"] == 'print("Hello friends")'
    assert len(exclusions) == 3 and data == original


def test_reading_cleanup_omits_dangling_example_without_removing_a_real_example():
    block = lambda text, kind="paragraph": {"text": text, "kind": kind, "source_chunk_ids": ["T0"]}
    data = {"title": "Lesson", "sections": [{"heading": "Examples", "blocks": [
        block("For example:"), block("print(30)", "code"), block("Example of a script:"), block("Output:")]}]}
    reading, excluded = clean_reading_document(data)
    assert len(reading["sections"][0]["blocks"]) == 2
    assert len(excluded) == 2


def test_reading_cleanup_removes_residual_audience_and_curriculum_boilerplate():
    text = "It is designed for individuals with varying experience. Many universities now use Python in the CBSE curriculum. Learning Python is easy and enjoyable. The function returns an integer."
    data = {"title": "Lesson", "sections": [{"heading": "Python", "blocks": [
        {"kind": "paragraph", "text": text, "source_chunk_ids": ["T0"]}]}]}
    reading, exclusions = clean_reading_document(data)
    assert reading["sections"][0]["blocks"][0]["text"] == "The function returns an integer."
    assert len(exclusions) == 3


@pytest.mark.parametrize("text", ["short", "α₂ .5 mW. " * 400, "longword " * 200])
def test_chunks_cover_every_character_and_bound_context(text):
    chunks = list(transcript_chunks(text, 500))
    assert "".join(c["text"] for c in chunks) == text
    assert all(len(c["text"]) <= 500 and len(c["context_before"]) <= 400 for c in chunks)
    for c in chunks:
        assert text[c["source_char_start"]:c["source_char_end"]] == c["text"]


@pytest.mark.parametrize("mode,error", [("empty", ValueError), ("invalid", ValueError),
                                         ("unavailable", OllamaError), ("truncated", OllamaError)])
def test_direct_failure_keeps_evidence_without_final_document(source_dir, mode, error):
    output = source_dir.parent / "failed-direct"
    with pytest.raises(error):
        generate_document(source_dir, output, WritingSettings(), DirectModel(mode))
    assert json.loads((output / "manifest.json").read_text())["state"] == "failed"
    assert (output / "model_calls").is_dir()
    assert not (output / "document.docx").exists()
    assert not (output / "reading_document.json").exists()


def test_direct_invalid_json_retries_once(source_dir):
    model = DirectModel("retry")
    result = generate_document(source_dir, source_dir.parent / "retry", WritingSettings(), model)
    assert result.is_file() and len(model.requests) == 2


def test_direct_budget_and_legacy_reuse_fail_before_calls(source_dir):
    model = DirectModel()
    with pytest.raises(ValueError, match="legacy"):
        generate_document(source_dir, source_dir.parent / "reuse", WritingSettings(), model, reuse_knowledge=source_dir)
    assert not model.requests
    with pytest.raises(ValueError, match="separate"):
        generate_document(source_dir, source_dir, WritingSettings(), model)


def test_direct_detects_source_changes_and_withdraws_final_export(source_dir):
    class Mutating(DirectModel):
        def chat(self, request):
            result = super().chat(request)
            (source_dir / "preserved.txt").write_text("changed")
            return result
    output = source_dir.parent / "mutated"
    with pytest.raises(ValueError, match="changed"):
        generate_document(source_dir, output, WritingSettings(), Mutating())
    assert not (output / "document.docx").exists()
    assert not json.loads((output / "integrity.json").read_text())["source_artifacts_unchanged"]


@pytest.mark.parametrize("mode", ["good", "unavailable"])
def test_default_api_direct_rewrite_without_audio_verification(tmp_path, audio_file, monkeypatch, mode):
    def fail(*args, **kwargs):
        pytest.fail("No second ASR pass is permitted on the default rewrite path")
    monkeypatch.setattr("videomind.worker.verify_run", fail)
    model = DirectModel(mode)
    app = create_app(Settings(model="primary-test"), IngestionSettings(storage_dir=tmp_path / "storage"),
                     MeasurementASR(), writer=model, start_worker=False)
    with TestClient(app, base_url="http://localhost") as client:
        response = client.post("/api/v1/jobs", data={"authorized": "true", "write_document": "true"},
                               files={"file": ("lesson.wav", audio_file.read_bytes(), "audio/wav")})
        assert response.status_code == 202
        location = response.headers["location"]
        app.state.worker.run_one()
        job = client.get(location).json()
        assert client.get(location + "/transcript").status_code == 200
        assert job["verification"] is None
        if mode == "unavailable":
            assert job["state"] == "failed" and job["error"]["code"] == "document_failed"
            assert client.get(location + "/document").status_code == 409
        else:
            assert job["state"] == "completed"
            assert "knowledge.json" not in job["artifacts"]
            assert "rewrite_audit.json" in job["artifacts"]
            word = client.get(location + "/document")
            assert word.status_code == 200 and Document(io.BytesIO(word.content))
            assert client.get(location + "/document?format=json").json()["pipeline"] == "direct"
            assert client.get(location + "/document?format=rewrite-audit").status_code == 200
            assert client.get(location + "/document?format=source-chunks").status_code == 200
            assert client.get(location + "/document?format=knowledge").status_code == 404
            assert client.get(location + "/document?format=../manifest").status_code == 422


def test_cli_defaults_to_direct_and_reports_unreviewed_status(source_dir, monkeypatch, capsys):
    from videomind import writing_cli
    monkeypatch.setattr("videomind.services.direct_rewrite_service.OllamaClient", lambda settings: DirectModel())
    code = writing_cli.main(["--source-dir", str(source_dir), "--output-dir", str(source_dir.parent / "cli-direct")])
    assert code == 2
    assert "document.docx" in capsys.readouterr().out
